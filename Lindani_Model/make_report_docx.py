"""
make_report_docx.py

Turns REPORT_DRAFT.md into REPORT_DRAFT.docx so the draft can be rewritten in
Word, which is where the report has to end up anyway.

This handles only the markdown actually used in the draft: headings, paragraphs,
bullet and numbered lists, tables, block quotes, and inline bold and code. It is
a convenience, not a general converter.

    python Lindani_Model/make_report_docx.py

Needs: pip install python-docx
"""

import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from split_check import LINDANI_DIR

SOURCE = LINDANI_DIR / os.environ.get("REPORT_SOURCE", "REPORT_DRAFT.md")
OUTPUT = SOURCE.with_suffix(".docx")

INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)")


def add_runs(paragraph, text):
    """Writes text into a paragraph, honouring **bold**, *italic* and `code`."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)


def split_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(document, lines):
    """A markdown table: header row, separator row, then body rows."""
    header = split_row(lines[0])
    body = [split_row(line) for line in lines[2:]]

    table = document.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, header):
        cell.paragraphs[0].text = ""
        add_runs(cell.paragraphs[0], text)
        for run in cell.paragraphs[0].runs:
            run.bold = True

    for row in body:
        cells = table.add_row().cells
        for cell, text in zip(cells, row[:len(header)]):
            cell.paragraphs[0].text = ""
            add_runs(cell.paragraphs[0], text)

    document.add_paragraph()


def main():
    lines = SOURCE.read_text(encoding="utf-8").split("\n")
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)

    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        # Tables: a pipe row followed by a separator row of dashes.
        if (stripped.startswith("|") and index + 1 < len(lines)
                and set(lines[index + 1].strip()) <= set("|-: ")
                and "-" in lines[index + 1]):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index])
                index += 1
            add_table(document, block)
            continue

        if stripped.startswith("#"):
            # Headings carry no inline formatting, so strip any markers rather
            # than printing them literally.
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip().replace("**", "").replace("`", "")
            document.add_heading(text, level=min(level - 1, 4))
        elif stripped.startswith("> "):
            paragraph = document.add_paragraph(style="Intense Quote")
            add_runs(paragraph, stripped[2:])
        elif stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_runs(paragraph, stripped[2:])
        elif re.match(r"^\d+\.\s", stripped):
            paragraph = document.add_paragraph(style="List Number")
            add_runs(paragraph, re.sub(r"^\d+\.\s", "", stripped))
        elif stripped.startswith("---"):
            document.add_paragraph()
        else:
            # Join wrapped lines into one paragraph.
            block = [stripped]
            index += 1
            while index < len(lines) and lines[index].strip() and not re.match(
                    r"^(#|-\s|\d+\.\s|\||>|---)", lines[index].strip()):
                block.append(lines[index].strip())
                index += 1
            paragraph = document.add_paragraph()
            add_runs(paragraph, " ".join(block))
            continue

        index += 1

    # A visible warning at the top, so a draft is never mistaken for the report.
    warning = document.paragraphs[0].insert_paragraph_before()
    run = warning.add_run(
        "DRAFT. Rewrite in your own words before submission. "
        "Items marked PENDING still need results that do not exist yet.")
    run.bold = True
    run.font.color.rgb = RGBColor(0xB4, 0x53, 0x2A)
    warning.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.save(OUTPUT)
    print(f"Wrote {OUTPUT}")
    print(f"{len(document.paragraphs)} paragraphs, {len(document.tables)} tables")


if __name__ == "__main__":
    main()
